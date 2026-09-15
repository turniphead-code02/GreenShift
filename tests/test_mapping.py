import openpyxl

def run_test():
    file_path = "1 CBAM SEE V2.1_Example Cement_final.xlsx"
    print(f"Loading {file_path}...")
    
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        print(f"Failed to load workbook: {e}")
        return
        
    print("Workbook loaded successfully.")
    
    # The dictionary from the plan
    CEMENT_CELL_MAP = {
        ("B_EmInst", "F14"): 1300,       # fuel_quantity (Coal)
        ("B_EmInst", "F15"): 10000,      # process_raw_qty (Clinker process)
        ("C_Emissions&Energy", "G15"): 1100, # electricity_mwh
        ("D_Processes", "L54"): 10000    # total_al_production (Clinker)
    }
    
    for (sheet_name, cell_ref), new_value in CEMENT_CELL_MAP.items():
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            old_value = ws[cell_ref].value
            ws[cell_ref] = new_value
            print(f"Updated {sheet_name}!{cell_ref} from {old_value} to {new_value}")
        else:
            print(f"Sheet {sheet_name} not found!")
            
    output_path = "Test_Output_Cement.xlsx"
    wb.save(output_path)
    print(f"Saved to {output_path}. Please open it in Excel and check Summary_Communication.")

if __name__ == "__main__":
    run_test()
